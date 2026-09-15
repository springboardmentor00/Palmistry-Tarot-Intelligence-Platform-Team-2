import io
from html import escape

import openpyxl
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


class ReportGenerator:
    @staticmethod
    def generate_pdf_report(reading_data: dict) -> io.BytesIO:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()
        title = ParagraphStyle("AetheriaTitle", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=colors.HexColor("#4c1d95"), spaceAfter=14)
        heading = ParagraphStyle("AetheriaHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#4c1d95"), spaceBefore=10, spaceAfter=6)
        body = ParagraphStyle("AetheriaBody", parent=styles["BodyText"], fontSize=9.5, leading=13, spaceAfter=8)
        small = ParagraphStyle("AetheriaSmall", parent=styles["BodyText"], fontSize=7.5, leading=10, textColor=colors.HexColor("#666666"))
        elements = [Paragraph("AETHERIA — SPIRITUAL INSIGHT REPORT", title)]
        metadata = [
            ["Reference", str(reading_data.get("id", "—")), "Date", str(reading_data.get("date", "—"))],
            ["Assessment", str(reading_data.get("type", "—")), "Focus", str(reading_data.get("focus", "—"))],
            ["Guidance score", str(reading_data.get("score", "—")), "Status", "Generated from persisted reading data"],
        ]
        table = Table(metadata, colWidths=[75, 180, 65, 180])
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f5f3ff")), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#ddd6fe")), ("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"), ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"), ("PADDING", (0, 0), (-1, -1), 7)]))
        elements += [table, Spacer(1, 16)]
        if reading_data.get("summary"):
            elements += [Paragraph("Executive Summary", heading), Paragraph(escape(str(reading_data["summary"])), body)]
        synthesis_text = escape(str(reading_data.get("synthesis", "No interpretation text was persisted for this report."))).replace("\n", "<br/>")
        elements += [Paragraph("Interpretation & Guidance", heading), Paragraph(synthesis_text, body)]
        cards = reading_data.get("tarot_cards")
        if isinstance(cards, list) and cards:
            elements.append(Paragraph("Tarot Draw", heading))
            card_rows = [["Position", "Card", "Orientation", "Meaning"]]
            card_rows.extend([[str(c.get("position", "—")), str(c.get("card_name", "—")), str(c.get("orientation", "—")), str(c.get("meaning", "—"))] for c in cards])
            card_table = Table(card_rows, colWidths=[85, 90, 70, 255], repeatRows=1)
            card_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4c1d95")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#ddd6fe")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("PADDING", (0, 0), (-1, -1), 5)]))
            elements += [card_table, Spacer(1, 10)]
        categories = reading_data.get("categories")
        if isinstance(categories, dict):
            elements.append(Paragraph("Insight Categories", heading))
            for key, value in categories.items():
                elements.append(Paragraph(escape(str(key).replace("_", " ").title()), ParagraphStyle("Cat", parent=heading, fontSize=10, spaceBefore=4)))
                elements.append(Paragraph(escape(str(value)).replace("\n", "<br/>"), body))
        elements += [Spacer(1, 20), Paragraph("Disclaimer", heading), Paragraph(escape(str(reading_data.get("disclaimer", "This report is intended for symbolic self-reflection and entertainment and is not professional advice."))), small)]
        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_excel_report(reading_data: dict) -> io.BytesIO:
        buffer = io.BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary"
        header_fill = PatternFill(fill_type="solid", fgColor="4C1D95")
        header_font = Font(bold=True, color="FFFFFF")
        ws["A1"] = "AETHERIA SPIRITUAL INSIGHT REPORT"
        ws["A1"].font = Font(bold=True, size=16, color="4C1D95")
        ws.merge_cells("A1:D1")
        rows = [("Reference", reading_data.get("id")), ("Date", reading_data.get("date")), ("Assessment", reading_data.get("type")), ("Focus", reading_data.get("focus")), ("Guidance score", reading_data.get("score")), ("Summary", reading_data.get("summary", "")), ("Interpretation", reading_data.get("synthesis", "")), ("Disclaimer", reading_data.get("disclaimer", ""))]
        for r, (key, value) in enumerate(rows, start=3):
            ws.cell(r, 1, key); ws.cell(r, 2, value); ws.cell(r, 1).font = Font(bold=True)
        ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 90
        categories = reading_data.get("categories")
        if isinstance(categories, dict):
            cat = wb.create_sheet("Insights"); cat.append(["Category", "Insight"])
            for c in cat[1]: c.fill = header_fill; c.font = header_font
            for key, value in categories.items(): cat.append([key, value])
            cat.column_dimensions["A"].width = 28; cat.column_dimensions["B"].width = 100
        cards = reading_data.get("tarot_cards")
        if isinstance(cards, list) and cards:
            tc = wb.create_sheet("Tarot Cards"); tc.append(["Position", "Card", "Orientation", "Meaning"])
            for c in tc[1]: c.fill = header_fill; c.font = header_font
            for card in cards: tc.append([card.get("position"), card.get("card_name"), card.get("orientation"), card.get("meaning")])
            tc.column_dimensions["A"].width = 28; tc.column_dimensions["B"].width = 28; tc.column_dimensions["C"].width = 18; tc.column_dimensions["D"].width = 90
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_analytics_excel_report(title: str, metrics: dict, daily_rows: list[dict]) -> io.BytesIO:
        buffer = io.BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Analytics Summary"
        header_fill = PatternFill(fill_type="solid", fgColor="4C1D95")
        header_font = Font(bold=True, color="FFFFFF")
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=16, color="4C1D95")
        ws.merge_cells("A1:D1")
        row = 3
        for key, value in metrics.items():
            ws.cell(row, 1, str(key).replace("_", " ").title())
            ws.cell(row, 2, value)
            ws.cell(row, 1).font = Font(bold=True)
            row += 1
        if daily_rows:
            daily = wb.create_sheet("Daily Metrics")
            headers = list(daily_rows[0].keys())
            daily.append(headers)
            for cell in daily[1]: cell.fill = header_fill; cell.font = header_font
            for item in daily_rows: daily.append([item.get(h) for h in headers])
            for col in range(1, len(headers) + 1): daily.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
        wb.save(buffer)
        buffer.seek(0)
        return buffer
