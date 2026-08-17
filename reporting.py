import io
from datetime import datetime

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
except ImportError:
    SimpleDocTemplate = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

class ReportGenerator:
    @staticmethod
    def generate_pdf_report(reading_data: dict) -> io.BytesIO:
        buffer = io.BytesIO()
        
        if SimpleDocTemplate is None:
            # Fallback simple text-based stream wrapped as file bytes
            text = f"""
            AETHERIA REPORT: {reading_data.get('type')}
            Ref ID: {reading_data.get('id')}
            Date: {reading_data.get('date')}
            Theme Focus: {reading_data.get('focus')}
            Guidance Score: {reading_data.get('score')}/100

            AI Synthesis Insight:
            {reading_data.get('synthesis')}

            Disclaimer: Readings are for spiritual self-reflection and entertainment.
            """
            buffer.write(text.encode('utf-8'))
            buffer.seek(0)
            return buffer

        # Build formal ReportLab PDF
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()
        
        # Custom styles
        title_style = ParagraphStyle(
            'MysticTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=22,
            textColor=colors.HexColor('#4c1d95'),
            spaceAfter=15
        )
        
        body_style = ParagraphStyle(
            'MysticBody',
            parent=styles['BodyText'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#333333'),
            spaceAfter=10
        )

        meta_style = ParagraphStyle(
            'MysticMeta',
            parent=styles['BodyText'],
            fontName='Helvetica-Bold',
            fontSize=11,
            textColor=colors.HexColor('#d97706')
        )

        elements = []
        
        # Title
        elements.append(Paragraph("AETHERIA SPIRITUAL INSIGHT REPORT", title_style))
        elements.append(Spacer(1, 10))
        
        # Metadata Table
        data = [
            [Paragraph(f"<b>Ref ID:</b> {reading_data.get('id')}", body_style), Paragraph(f"<b>Date:</b> {reading_data.get('date')}", body_style)],
            [Paragraph(f"<b>Assessment:</b> {reading_data.get('type')}", body_style), Paragraph(f"<b>Theme Focus:</b> {reading_data.get('focus')}", body_style)],
            [Paragraph(f"<b>Guidance Score:</b> {reading_data.get('score')}/100", meta_style), ""]
        ]
        t = Table(data, colWidths=[250, 250])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f5f3ff')),
            ('PADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,2), (-1,2), 12),
            ('SPAN', (0,2), (1,2)),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 20))
        
        # Detailed Insight
        elements.append(Paragraph("<b>AI Synthesis & Guidance</b>", ParagraphStyle('Heading2', parent=styles['Heading2'], textColor=colors.HexColor('#4c1d95'))))
        elements.append(Spacer(1, 5))
        elements.append(Paragraph(reading_data.get('synthesis', ''), body_style))
        elements.append(Spacer(1, 20))

        # Disclaimer
        elements.append(Spacer(1, 40))
        disclaimer_text = "<b>Disclaimer:</b> This report is generated symbolically using computational models. It does not constitute medical, legal, or professional counseling advice."
        elements.append(Paragraph(disclaimer_text, ParagraphStyle('Disclaimer', parent=styles['Italic'], fontSize=8, textColor=colors.HexColor('#777777'))))

        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_excel_report(reading_data: dict) -> io.BytesIO:
        buffer = io.BytesIO()
        
        if openpyxl is None:
            # Fallback CSV format
            csv = f"ID,Date,Type,Theme,Guidance Score\n{reading_data.get('id')},{reading_data.get('date')},{reading_data.get('type')},{reading_data.get('focus')},{reading_data.get('score')}"
            buffer.write(csv.encode('utf-8'))
            buffer.seek(0)
            return buffer

        # Build openpyxl Excel file
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Aetheria Assessment"

        ws['A1'] = "AETHERIA SPIRITUAL INSIGHT DATA"
        ws['A1'].font = openpyxl.styles.Font(bold=True, size=14, color="4C1D95")
        ws.merge_cells('A1:E1')

        headers = ["Ref ID", "Date", "Reading Type", "Theme Focus", "Guidance Score"]
        ws.append([])  # Spacer row
        ws.append(headers)
        
        # Apply header formatting
        for col_num in range(1, 6):
            cell = ws.cell(row=3, column=col_num)
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")

        data_row = [
            reading_data.get('id'),
            reading_data.get('date'),
            reading_data.get('type'),
            reading_data.get('focus'),
            reading_data.get('score')
        ]
        ws.append(data_row)
        
        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        wb.save(buffer)
        buffer.seek(0)
        return buffer
